"""Create a shareable Word edition from the authoritative Chinese Markdown manual.

Run with a Python environment containing python-docx. This does not access the network.
"""

import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
source = ROOT / "docs" / "功能与部署文档.md"
document = Document()
section = document.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = section.right_margin = Cm(2.3)
section.page_width, section.page_height = Cm(21), Cm(29.7)
styles = document.styles
for name in ["Normal", "Body Text", "List Bullet", "List Number"]:
    styles[name].font.name = "Microsoft YaHei"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles[name].font.size = Pt(10)
    styles[name].font.color.rgb = RGBColor.from_string("344C36")
    styles[name].paragraph_format.space_after = Pt(7)
    styles[name].paragraph_format.line_spacing = 1.3
for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
    styles[name].font.name = "Microsoft YaHei"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles[name].font.color.rgb = RGBColor.from_string("244F39")
header = section.header.paragraphs[0]
header.text = "TEACHAGENT     /     学校本地课堂洞察与教学成长工作台"
header.style = styles["Caption"]
footer = section.footer.paragraphs[0]
footer.text = (
    "TeachAgent 1.1.0    ·    功能与部署文档                                      "
)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


def clean(text):
    return re.sub(r"\*\*(.*?)\*\*|`([^`]+)`", lambda m: m[1] or m[2], text).strip()


lines = source.read_text(encoding="utf-8").splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if line.startswith("```"):
        language = line[3:]
        code = []
        i += 1
        while i < len(lines) and not lines[i].startswith("```"):
            code.append(lines[i])
            i += 1
        if language == "mermaid":
            document.add_paragraph(
                "教师浏览器 → 学校内网 FastAPI → 本地数据库 / 按教师隔离的文件目录\n任务队列 → FFmpeg / PyAV → Whisper → 文本蒸馏与报告\n教研追问 → 本机 Ollama / 本地规则 / 显式启用的 OpenAI 兼容服务（仅文本）"
            )
        else:
            paragraph = document.add_paragraph()
            run = paragraph.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "F0F5EB")
            paragraph._p.get_or_add_pPr().append(shade)
    elif line.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [clean(c) for c in lines[i].strip("|").split("|")]
            if not all(re.fullmatch(r"[-: ]+", c) for c in cells):
                rows.append(cells)
            i += 1
        table = document.add_table(rows=0, cols=len(rows[0]))
        table.style = "Light Shading Accent 1"
        for index, row in enumerate(rows):
            cells = table.add_row().cells
            for c, value in zip(cells, row):
                c.text = value
                for paragraph in c.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8)
            trpr = table.rows[-1]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if index == 0:
                repeat = OxmlElement("w:tblHeader")
                trpr.append(repeat)
        document.add_paragraph()
        continue
    elif line.startswith("# "):
        document.add_paragraph("TeachAgent", "Title")
        document.add_paragraph("功能与部署文档", "Subtitle")
        document.add_paragraph("让课堂被看见，让成长有迹可循。")
        image = ROOT / "docs" / "images" / "dashboard.png"
        if image.exists():
            document.add_picture(str(image), width=Cm(15.8))
        document.add_page_break()
    elif line.startswith("## "):
        document.add_heading(clean(line[3:]), level=1)
    elif line.startswith("### "):
        document.add_heading(clean(line[4:]), level=2)
    elif line.startswith("- "):
        document.add_paragraph(clean(line[2:]), "List Bullet")
    elif re.match(r"^\d+\. ", line):
        document.add_paragraph(clean(re.sub(r"^\d+\. ", "", line)), "List Number")
    elif line.strip():
        document.add_paragraph(clean(line))
    i += 1
output = ROOT / "docs" / "TeachAgent-功能与部署文档.docx"
document.save(output)
print(
    f"Created {output.name}: {len(document.paragraphs)} paragraphs, {len(document.tables)} tables."
)
