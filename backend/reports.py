import html
import io
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def word_report(lesson, analysis):
    """Generate an actual editable OOXML report, without sending data anywhere."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    for name, size in [
        ("Normal", 11),
        ("Title", 24),
        ("Subtitle", 14),
        ("Heading 1", 16),
        ("Heading 2", 13),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("244E3B")
        fonts = style.element.get_or_add_rPr().rFonts
        fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        for attr in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme"]:
            fonts.attrib.pop(qn("w:" + attr), None)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.space_before = Pt(
            18 if name == "Heading 1" else 14 if name == "Heading 2" else 0
        )
        style.paragraph_format.space_after = Pt(
            10 if name == "Heading 1" else 7 if name == "Heading 2" else 6
        )
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)

    def clean(value):
        return re.sub(
            r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", str(value)
        )

    def paragraph(text, style=None):
        return doc.add_paragraph(clean(text), style)

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        widths = [9360 // len(headers)] * len(headers)
        props = t._tbl.tblPr
        props.find(qn("w:tblW")).set(qn("w:w"), "9360")
        props.find(qn("w:tblW")).set(qn("w:type"), "dxa")
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), "120")
        indent.set(qn("w:type"), "dxa")
        props.append(indent)
        margins = OxmlElement("w:tblCellMar")
        for edge, value in [("top", 80), ("bottom", 80), ("start", 120), ("end", 120)]:
            el = OxmlElement("w:" + edge)
            el.set(qn("w:w"), str(value))
            el.set(qn("w:type"), "dxa")
            margins.append(el)
        props.append(margins)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for i, name in enumerate(headers):
            t.rows[0].cells[i].text = name
            t.rows[0].cells[i].paragraphs[0].paragraph_format.keep_with_next = True
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E8EEF5")
            t.rows[0].cells[i]._tc.get_or_add_tcPr().append(shade)
        for row in rows:
            for cell, value in zip(t.add_row().cells, row):
                cell.text = clean(value)
        for col, width in zip(t._tbl.tblGrid, widths):
            col.set(qn("w:w"), str(width))
        for row in t.rows:
            for cell, width in zip(row.cells, widths):
                cell._tc.get_or_add_tcPr().find(qn("w:tcW")).set(qn("w:w"), str(width))

    doc.core_properties.title = clean(lesson["title"] + " · 教学分析报告")
    doc.core_properties.author = "TeachAgent"
    section.header.paragraphs[0].text = "TeachAgent · 课堂教学分析"
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    paragraph(lesson["title"], "Title")
    paragraph("教学分析报告", "Subtitle")
    paragraph(f"学科：{lesson['subject']}    班级：{lesson['class_name']}")
    paragraph(
        "数据性质："
        + ("合成示例，仅用于体验。" if lesson["example"] else "课堂转写分析。")
    )
    paragraph("授课方式画像", "Heading 1")
    paragraph(analysis["summary"])
    labels = {
        "characters": "有效字符",
        "segments": "话语片段",
        "questions": "提问线索",
        "interactions": "互动邀请线索",
        "speech_rate": "话语语速（字/分钟）",
        "labeled_student_segments": "显式学生标签片段",
    }
    table(
        ["指标", "结果"],
        [(label, analysis["metrics"][key]) for key, label in labels.items()],
    )
    paragraph("环节时间分配", "Heading 1")
    table(
        ["环节", "有效话语时长（秒）", "占比"],
        [
            (s["name"], s["seconds"], f"{s['percent']}%")
            for s in analysis["distribution"]
        ],
    )
    paragraph("可优化点清单", "Heading 1")
    for suggestion in analysis["suggestions"]:
        paragraph(suggestion["title"], "Heading 2")
        evidence = paragraph(
            f"证据（片段 #{suggestion['segment_id']}）：{suggestion['evidence']}"
        )
        evidence.paragraph_format.keep_with_next = True
        paragraph("建议：" + suggestion["action"])
    paragraph("方法与边界", "Heading 1")
    for text in analysis["limitations"]:
        paragraph(text)
    paragraph("转写原文", "Heading 1")
    for segment in analysis["segments"]:
        start = segment["start"]
        paragraph(
            f"[{int(start // 60):02d}:{int(start % 60):02d}] #{segment['id']} 【{segment['stage']}】{segment['text']}"
        )
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def markdown(lesson, analysis):
    lines = [
        f"# {lesson['title']} · 教学分析报告",
        "",
        f"学科：{lesson['subject']}　班级：{lesson['class_name']}",
        "",
        "> 数据性质："
        + ("合成示例，仅用于体验。" if lesson["example"] else "课堂转写与教学分析；转写来源及时间精度请见方法边界。"),
        "",
        "## 授课方式画像",
        "",
        analysis["summary"],
        "",
        "| 指标 | 结果 |",
        "| --- | --- |",
    ]
    labels = {
        "characters": "有效字符",
        "segments": "话语片段",
        "questions": "提问线索",
        "interactions": "互动邀请线索",
        "speech_rate": "话语语速（字/分钟）",
        "labeled_student_segments": "显式学生标签片段",
    }
    lines += [
        f"| {label} | {analysis['metrics'][key]} |" for key, label in labels.items()
    ]
    lines += ["", "## 环节时间分配", ""]
    lines += [
        f"- {s['name']}：{s['seconds']} 秒，占有效话语时长 {s['percent']}%"
        for s in analysis["distribution"]
    ]
    lines += ["", "## 可优化点清单", ""]
    for s in analysis["suggestions"]:
        lines += [
            f"### {s['title']}",
            "",
            f"证据（片段 #{s['segment_id']}）：{s['evidence']}",
            "",
            f"建议：{s['action']}",
            "",
        ]
    lines += ["## 方法与边界", ""] + ["- " + text for text in analysis["limitations"]]
    lines += ["", "## 转写原文", ""]
    for s in analysis["segments"]:
        stamp = f"{int(s['start'] // 60):02d}:{int(s['start'] % 60):02d}"
        lines.append(f"- [{stamp}] #{s['id']} 【{s['stage']}】{s['text']}")
    return "\n".join(lines)


def html_report(lesson, analysis):
    content = html.escape(markdown(lesson, analysis))
    return f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>TeachAgent 教学报告</title><style>body{{font-family:system-ui,sans-serif;max-width:960px;margin:48px auto;padding:24px;color:#163e36;line-height:1.9}}pre{{white-space:pre-wrap;font:inherit}}@media print{{body{{margin:0}}}}</style><h1>TeachAgent · 教学分析报告</h1><p>可通过浏览器打印菜单保存为 PDF。学校本地数据，请妥善保管。</p><pre>{content}</pre></html>'
