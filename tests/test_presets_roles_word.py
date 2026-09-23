import io
import json
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from backend import model_service
from backend.db import connect, init_db
from backend.main import app


def test_presets_key_only_switching_encryption_and_account_isolation(clients):
    c, ids = clients
    alice = c["alice"]
    for preset_id, key in [
        ("deepseek", "fixture-deepseek-key"),
        ("openai", "fixture-openai-key"),
    ]:
        result = alice.put(
            "/api/model-settings",
            json={
                "provider": "openai",
                "preset_id": preset_id,
                "api_key": key,
                "allow_external": True,
            },
        )
        assert result.status_code == 200, result.text
        assert result.json()["model"] == model_service.PRESETS[preset_id]["model"]
        assert result.json()["base_url"] == model_service.PRESETS[preset_id]["base_url"]
        assert key not in result.text
    # Key-only save is enough; even forged endpoint fields cannot redirect a preset key.
    result = alice.put(
        "/api/model-settings",
        json={
            "provider": "openai",
            "preset_id": "deepseek",
            "base_url": "http://127.0.0.1:1",
            "model": "forged",
            "allow_external": True,
        },
    )
    assert result.status_code == 200
    settings, encrypted = model_service.runtime_settings(ids["alice"])
    _, url, headers, payload, _ = model_service._request(settings, encrypted, [])
    assert url == "https://api.deepseek.com/v1/chat/completions"
    assert headers["Authorization"] == "Bearer fixture-deepseek-key"
    assert payload["model"] == "deepseek-chat"
    with connect() as db:
        raw = str([tuple(r) for r in db.execute("SELECT * FROM model_profiles")])
    assert "fixture-deepseek-key" not in raw and "fixture-openai-key" not in raw
    profiles = alice.get("/api/model-settings/profiles")
    assert len(profiles.json()) == 2 and "fixture-" not in profiles.text
    assert c["bob"].get("/api/model-settings/profiles").json() == []
    assert not c["bob"].get("/api/model-settings").json()["api_key_configured"]
    clear = alice.put(
        "/api/model-settings",
        json={
            "provider": "openai",
            "preset_id": "deepseek",
            "clear_api_key": True,
            "allow_external": True,
        },
    )
    assert not clear.json()["api_key_configured"]
    saved = {
        p["preset_id"]: p for p in alice.get("/api/model-settings/profiles").json()
    }
    assert (
        not saved["deepseek"]["api_key_configured"]
        and saved["openai"]["api_key_configured"]
    )
    assert (
        alice.put(
            "/api/model-settings", json={"provider": "openai", "preset_id": "unknown"}
        ).status_code
        == 422
    )
    assert (
        alice.put(
            "/api/model-settings", json={"provider": "openai", "preset_id": "deepseek"}
        ).status_code
        == 400
    )


def test_superadmin_migration_and_admin_creation_boundary(clients):
    c, ids = clients
    assert c["admin"].get("/api/auth/me").json()["user"]["is_superadmin"]
    # Idempotent startup retains existing identity, password and data.
    init_db()
    assert c["admin"].get("/api/auth/me").status_code == 200
    data = {
        "username": "school-admin",
        "name": "学校管理员",
        "password": "Fixture-password-123",
        "role": "admin",
    }
    result = c["admin"].post("/api/admin/users", json=data)
    assert result.status_code == 201
    with TestClient(app) as ordinary:
        login = ordinary.post(
            "/api/auth/login",
            json={"username": data["username"], "password": data["password"]},
        )
        assert not login.json()["user"]["is_superadmin"]
        ordinary.headers["x-csrf-token"] = login.json()["csrf"]
        assert (
            ordinary.post(
                "/api/admin/users", json={**data, "username": "other-admin"}
            ).status_code
            == 403
        )
        assert (
            ordinary.post(
                "/api/admin/users",
                json={**data, "username": "new-teacher", "role": "teacher"},
            ).status_code
            == 201
        )
        assert (
            ordinary.get(f"/api/admin/users/{ids['admin']}/backup").status_code == 403
        )
        assert (
            ordinary.get(f"/api/admin/users/{ids['alice']}/backup").status_code == 200
        )
        assert (
            ordinary.get(f"/api/admin/users/{result.json()['id']}/backup").status_code
            == 200
        )
    assert c["alice"].post("/api/admin/users", json=data).status_code == 403
    assert (
        c["admin"]
        .post("/api/admin/users", json={**data, "username": "Admin"})
        .status_code
        == 409
    )
    assert (
        c["admin"]
        .post(
            "/api/admin/users", json={**data, "username": "super", "role": "superadmin"}
        )
        .status_code
        == 400
    )


def test_word_download_is_real_docx_complete_and_isolated(clients):
    c, _ = clients
    assert c["alice"].get("/api/health").json()["version"] == app.version
    lesson = (
        c["alice"]
        .post(
            "/api/lessons/text",
            json={
                "title": "中文课堂 <示例>",
                "subject": "数学",
                "class_name": "一班",
                "text": "教师：今天我们学习函数。\n学生：为什么会变化？\n教师：请练习并总结。",
            },
        )
        .json()
    )
    url = f"/api/lessons/{lesson['id']}/report"
    for suffix in ["", "?format=docx"]:
        report = c["alice"].get(url + suffix)
        assert report.status_code == 200
        assert (
            report.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert report.headers["content-disposition"].endswith('.docx"')
        assert report.content.startswith(b"PK")
        doc = Document(io.BytesIO(report.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert lesson["title"] in text
        for heading in [
            "授课方式画像",
            "环节时间分配",
            "可优化点清单",
            "方法与边界",
            "转写原文",
        ]:
            assert heading in text
        for segment in lesson["analysis"]["segments"]:
            assert segment["text"] in text
        assert len(doc.tables) == 2
    assert c["bob"].get(url).status_code == 404
    assert c["admin"].get(url).status_code == 404


def test_pages_word_matches_synthetic_lesson():
    root = Path(__file__).resolve().parents[1]
    demo = json.loads((root / "frontend/src/demo.json").read_text(encoding="utf8"))
    doc = Document(root / "frontend/public/example-report.docx")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert demo["title"] in text and "合成示例" in text
    for segment in demo["analysis"]["segments"]:
        assert segment["text"] in text
